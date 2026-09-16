# Copyright (c) 2026, osher and contributors
# For license information, please see license.txt

"""Monthly consolidated aircraft Excel report (2 sheets, RTL Hebrew, all totals = formulas).

Generates ONE xlsx per aircraft per month from submitted Work Reports
(docstatus=1, filtered by work_date month):
  Sheet 1 "עבודות וזמנים" - one row per Work Report:
      תאריך | מס' דוח | טכנאי | שעת התחלה | שעת סיום | שעות (עשרוני) |
      שעות (HH:MM) | סוגי עבודה | הערות
      Summary row: =COUNTA(report-col) for report count, =SUM(decimal-col)
      for total hours.
  Sheet 2 "חלפים מסכמים" - parts aggregated by part_number across the month:
      מק"ט | שם פריט | כמות סה"כ | מחיר יחידה | עלות סה"כ (=qty*price) | נצרך בדוחות
      Summary row: =SUM(qty-col), =SUM(cost-col).

Both sheets: RTL, navy headers, zebra striping, autofilter, freeze panes,
fullCalcOnLoad so Excel recalculates on open.
"""

import calendar
import datetime
from io import BytesIO

import frappe
from frappe.utils import add_months, getdate, today


NAVY = "1F3864"
ZEBRA = "EDF2F9"
SUM_BG = "D9EAD3"

HEBREW_MONTHS = {
    1: "ינואר",
    2: "פברואר",
    3: "מרץ",
    4: "אפריל",
    5: "מאי",
    6: "יוני",
    7: "יולי",
    8: "אוגוסט",
    9: "ספטמבר",
    10: "אוקטובר",
    11: "נובמבר",
    12: "דצמבר",
}


def _work_type_labels(report):
    """Hebrew work-type labels for one Work Report doc (same labels as per-report email)."""
    labels = []
    if report.get("oxygen_fill"):
        labels.append("מילוי חמצן")
    if report.get("nitrogen_fill"):
        labels.append("מילוי חנקן")
    if report.get("tks_fill"):
        labels.append("מילוי TKS")
    if report.get("hydraulic_oil_fill"):
        labels.append("מילוי שמן הידראולי")
    if report.get("gpu_usage"):
        gpu_hours = report.get("gpu_hours") or 0
        labels.append(f"GPU ({gpu_hours} שעות)" if gpu_hours else "GPU")
    if report.get("wi_di"):
        labels.append("WI/DI")
    return ", ".join(labels)


def _strip_html(text):
    """Notes are stored as HTML (Text Editor) - strip tags for the Excel cell."""
    if not text:
        return ""
    import re

    text = re.sub(r"<br\s*/?>", " ", text)
    text = re.sub(r"</(p|div|li|tr|h[1-6])>", " ", text)
    text = re.sub(r"<[^>]+>", "", text)
    text = text.replace("&nbsp;", " ").replace("&amp;", "&")
    return " ".join(text.split())


def _fmt_time(value):
    """Format a Frappe Time value (timedelta) as HH:MM string."""
    if value is None:
        return ""
    if isinstance(value, datetime.timedelta):
        total_seconds = int(value.total_seconds())
        return f"{total_seconds // 3600:02d}:{(total_seconds % 3600) // 60:02d}"
    text = str(value)
    parts = text.split(":")
    if len(parts) >= 2:
        return f"{parts[0].zfill(2)}:{parts[1].zfill(2)}"
    return text


def _get_reports(aircraft, year, month):
    """Submitted Work Reports for one aircraft in one calendar month."""
    start = datetime.date(year, month, 1)
    end = datetime.date(year, month, calendar.monthrange(year, month)[1])
    return frappe.get_all(
        "Work Report",
        filters={
            "aircraft": aircraft,
            "work_date": ["between", [str(start), str(end)]],
            "docstatus": 1,
        },
        fields=[
            "name",
            "work_date",
            "technician",
            "start_time",
            "end_time",
            "total_hours",
            "total_hours_display",
            "oxygen_fill",
            "nitrogen_fill",
            "tks_fill",
            "hydraulic_oil_fill",
            "gpu_usage",
            "gpu_hours",
            "wi_di",
            "notes",
        ],
        order_by="work_date asc, start_time asc, name asc",
    )


def _get_month_parts(report_names):
    """Part rows for the given reports, with item names from Inventory Item."""
    if not report_names:
        return []
    rows = frappe.get_all(
        "Work Report Part",
        filters={"parent": ["in", report_names]},
        fields=["parent", "part", "part_number", "quantity_used", "unit_price", "line_total"],
        order_by="parent asc, idx asc",
    )
    item_names = {}
    item_ids = {r.get("part") for r in rows if r.get("part")}
    if item_ids:
        for item in frappe.get_all(
            "Inventory Item",
            filters={"name": ["in", list(item_ids)]},
            fields=["name", "item_name"],
        ):
            item_names[item["name"]] = item.get("item_name") or ""
    for r in rows:
        r["item_name"] = item_names.get(r.get("part"), "")
    return rows


def get_monthly_report_data(aircraft, year, month):
    """Collect works rows + aggregated parts for one aircraft/month (testable, no xlsx)."""
    reports = _get_reports(aircraft, year, month)
    technician_names = {}
    for rep in reports:
        tech = rep.get("technician")
        if tech and tech not in technician_names:
            technician_names[tech] = frappe.db.get_value("User", tech, "full_name") or tech

    works = []
    for rep in reports:
        works.append(
            {
                "work_date": getdate(rep.get("work_date")).isoformat() if rep.get("work_date") else "",
                "report": rep.get("name"),
                "technician": technician_names.get(rep.get("technician"), rep.get("technician") or ""),
                "start": _fmt_time(rep.get("start_time")),
                "end": _fmt_time(rep.get("end_time")),
                "hours_decimal": float(rep.get("total_hours") or 0),
                "hours_hhmm": rep.get("total_hours_display") or "",
                "work_types": _work_type_labels(rep),
                "notes": _strip_html(rep.get("notes")),
            }
        )

    parts = _get_month_parts([r["name"] for r in reports])
    aggregated = {}
    for p in parts:
        key = p.get("part_number") or p.get("part") or "N/A"
        entry = aggregated.setdefault(
            key,
            {
                "part_number": key,
                "item_name": p.get("item_name") or "",
                "qty": 0.0,
                "unit_price": float(p.get("unit_price") or 0),
                "reports": [],
            },
        )
        entry["qty"] += float(p.get("quantity_used") or 0)
        if not entry["item_name"] and p.get("item_name"):
            entry["item_name"] = p.get("item_name")
        if p.get("unit_price") and not entry["unit_price"]:
            entry["unit_price"] = float(p.get("unit_price"))
        if p.get("parent") not in entry["reports"]:
            entry["reports"].append(p.get("parent"))

    parts_rows = [
        {
            "part_number": e["part_number"],
            "item_name": e["item_name"],
            "qty": e["qty"],
            "unit_price": e["unit_price"],
            "reports": ", ".join(sorted(e["reports"])),
        }
        for e in sorted(aggregated.values(), key=lambda e: e["part_number"])
    ]
    return {"works": works, "parts": parts_rows}


def generate_monthly_aircraft_report(aircraft, year, month):
    """Build the monthly xlsx for one aircraft. Returns bytes.

    Throws if no submitted reports found for the aircraft/month.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    data = get_monthly_report_data(aircraft, year, month)
    if not data["works"]:
        frappe.throw(f"לא נמצאו דוחות מאושרים למטוס {aircraft} לחודש {month:02d}/{year}")

    month_he = HEBREW_MONTHS.get(int(month), str(month))
    title_suffix = f"מטוס {aircraft} — {month_he} {year}"

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=NAVY)
    title_font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    title_fill = PatternFill("solid", fgColor=NAVY)
    body_font = Font(name="Arial", size=11)
    bold_font = Font(name="Arial", size=11, bold=True)
    sum_font = Font(name="Arial", size=12, bold=True)
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center", wrap_text=True)
    zebra_fill = PatternFill("solid", fgColor=ZEBRA)
    sum_fill = PatternFill("solid", fgColor=SUM_BG)
    subtitle_font = Font(name="Arial", size=10, italic=True, color="404040")

    wb = Workbook()

    # ============ SHEET 1: works ============
    ws = wb.active
    ws.title = "עבודות וזמנים"
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.merge_cells("A1:I1")
    c = ws["A1"]
    c.value = f"דוח עבודות חודשי — {title_suffix}"
    c.font = title_font
    c.fill = title_fill
    c.alignment = center
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:I2")
    c2 = ws["A2"]
    c2.value = "מופק ממערכת Arrow Aviation MX"
    c2.font = subtitle_font
    c2.alignment = center
    ws.row_dimensions[2].height = 20

    headers = [
        "תאריך",
        "מס' דוח",
        "טכנאי",
        "שעת התחלה",
        "שעת סיום",
        "שעות (עשרוני)",
        "שעות (HH:MM)",
        "סוגי עבודה",
        "הערות",
    ]
    hr = 3
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=hr, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[hr].height = 28

    r0 = hr + 1
    for i, w in enumerate(data["works"]):
        rr = r0 + i
        ws.row_dimensions[rr].height = 24
        vals = [
            w["work_date"],
            w["report"],
            w["technician"],
            w["start"],
            w["end"],
            w["hours_decimal"],
            w["hours_hhmm"],
            w["work_types"],
            w["notes"],
        ]
        for col, v in enumerate(vals, start=1):
            cell = ws.cell(row=rr, column=col, value=v)
            cell.font = body_font
            cell.border = border
            cell.alignment = center if col in (1, 2, 3, 4, 5, 7) else right
            if i % 2 == 1:
                cell.fill = zebra_fill
        ws.cell(row=rr, column=6).number_format = "0.00"

    n = len(data["works"])
    srow = r0 + n
    ws.row_dimensions[srow].height = 28
    ws.merge_cells(start_row=srow, start_column=1, end_row=srow, end_column=4)
    for col in range(1, 5):
        cell = ws.cell(row=srow, column=col)
        cell.fill = sum_fill
        cell.border = border
    lab = ws.cell(row=srow, column=1, value="סה״כ חודש")
    lab.font = sum_font
    lab.alignment = center
    cc = ws.cell(row=srow, column=5, value=f"=COUNTA(B{r0}:B{r0 + n - 1})")
    cc.font = sum_font
    cc.alignment = center
    cc.border = border
    cc.fill = sum_fill
    sc = ws.cell(row=srow, column=6, value=f"=SUM(F{r0}:F{r0 + n - 1})")
    sc.font = sum_font
    sc.alignment = center
    sc.border = border
    sc.fill = sum_fill
    sc.number_format = "0.00"
    nc = ws.cell(row=srow, column=7, value="← מחושב אוטומטית")
    nc.font = Font(name="Arial", size=10, italic=True)
    nc.alignment = center
    nc.fill = sum_fill
    nc.border = border
    for col in (8, 9):
        x = ws.cell(row=srow, column=col)
        x.fill = sum_fill
        x.border = border

    for i, w in enumerate([12, 16, 14, 12, 12, 13, 13, 42, 28], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A{hr}:I{r0 + n - 1}"

    # ============ SHEET 2: parts ============
    ws2 = wb.create_sheet("חלפים מסכמים")
    ws2.sheet_view.rightToLeft = True
    ws2.sheet_properties.pageSetUpPr.fitToPage = True

    ws2.merge_cells("A1:F1")
    t = ws2["A1"]
    t.value = f"סיכום חלפים חודשי — {title_suffix}"
    t.font = title_font
    t.fill = title_fill
    t.alignment = center
    ws2.row_dimensions[1].height = 30

    ws2.merge_cells("A2:F2")
    t2 = ws2["A2"]
    t2.value = "מקובץ לפי מק״ט — כמויות מסוכמות מכל דוחות החודש"
    t2.font = subtitle_font
    t2.alignment = center
    ws2.row_dimensions[2].height = 20

    h2 = ["מק״ט", "שם פריט", "כמות סה״כ", "מחיר יחידה (₪)", "עלות סה״כ (₪)", "נצרך בדוחות"]
    hr2 = 3
    for col, h in enumerate(h2, start=1):
        cell = ws2.cell(row=hr2, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    ws2.row_dimensions[hr2].height = 28

    p0 = hr2 + 1
    for i, p in enumerate(data["parts"]):
        rr = p0 + i
        ws2.row_dimensions[rr].height = 24
        vals = [p["part_number"], p["item_name"], p["qty"], p["unit_price"], f"=C{rr}*D{rr}", p["reports"]]
        for col, v in enumerate(vals, start=1):
            cell = ws2.cell(row=rr, column=col, value=v)
            cell.font = body_font
            cell.border = border
            cell.alignment = center if col in (1, 3, 4, 5) else right
            if i % 2 == 1:
                cell.fill = zebra_fill
        ws2.cell(row=rr, column=3).number_format = "0"
        ws2.cell(row=rr, column=4).number_format = "#,##0.00"
        ws2.cell(row=rr, column=5).number_format = "#,##0.00"

    m = len(data["parts"])
    if m:
        sr2 = p0 + m
        ws2.row_dimensions[sr2].height = 28
        ws2.merge_cells(start_row=sr2, start_column=1, end_row=sr2, end_column=2)
        lab2 = ws2.cell(row=sr2, column=1, value="סה״כ חודש")
        lab2.font = sum_font
        lab2.alignment = center
        for col in range(1, 7):
            ws2.cell(row=sr2, column=col).fill = sum_fill
            ws2.cell(row=sr2, column=col).border = border
            ws2.cell(row=sr2, column=col).font = sum_font
        qc = ws2.cell(row=sr2, column=3, value=f"=SUM(C{p0}:C{p0 + m - 1})")
        qc.alignment = center
        qc.number_format = "0"
        tc = ws2.cell(row=sr2, column=5, value=f"=SUM(E{p0}:E{p0 + m - 1})")
        tc.alignment = center
        tc.number_format = "#,##0.00"
        ws2.auto_filter.ref = f"A{hr2}:F{p0 + m - 1}"
    else:
        ws2.cell(row=p0, column=1, value="לא נצרכו חלקים בחודש זה").font = bold_font
        ws2.merge_cells(start_row=p0, start_column=1, end_row=p0, end_column=6)
        ws2.auto_filter.ref = f"A{hr2}:F{hr2}"

    for i, w in enumerate([18, 34, 12, 16, 16, 40], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws2.freeze_panes = "A4"

    wb.properties.creator = "Arrow Aviation MX"
    wb.properties.title = f"דוח חודשי {aircraft} {month:02d}-{year}"
    wb.calculation.fullCalcOnLoad = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _get_monthly_recipients():
    """Monthly recipients: monthly_report_recipients, fallback to work_report_recipients."""
    try:
        settings = frappe.get_doc("ARROW MX Settings")
    except Exception:
        return []
    recipients = []
    for row in getattr(settings, "monthly_report_recipients", None) or []:
        if row.get("email"):
            recipients.append(row.get("email"))
    if not recipients:
        for row in getattr(settings, "work_report_recipients", None) or []:
            if row.get("email"):
                recipients.append(row.get("email"))
    return recipients


def _resolve_month(year=None, month=None):
    if year and month:
        return int(year), int(month)
    d = getdate(add_months(today(), -1))
    return d.year, d.month


def send_monthly_reports(year=None, month=None):
    """Send ONE consolidated xlsx email per aircraft for the given (or previous) month.

    Returns {"year": y, "month": m, "aircraft": [...], "skipped": [...]}.
    Queues via frappe.sendmail (Email Queue) - no direct external send.
    """
    year, month = _resolve_month(year, month)
    start = datetime.date(year, month, 1)
    end = datetime.date(year, month, calendar.monthrange(year, month)[1])

    aircraft_list = frappe.get_all(
        "Work Report",
        filters={"work_date": ["between", [str(start), str(end)]], "docstatus": 1},
        fields=["aircraft"],
    )
    aircrafts = sorted({r["aircraft"] for r in aircraft_list if r.get("aircraft")})

    recipients = _get_monthly_recipients()
    if not recipients:
        frappe.log_error("Monthly report: no recipients configured", "Monthly Report")
        return {"year": year, "month": month, "aircraft": [], "skipped": aircrafts}

    month_he = HEBREW_MONTHS.get(int(month), str(month))
    sent, skipped = [], []
    for aircraft in aircrafts:
        try:
            content = generate_monthly_aircraft_report(aircraft, year, month)
        except Exception as e:
            frappe.log_error(f"Monthly report build failed for {aircraft}: {e}", "Monthly Report")
            skipped.append(aircraft)
            continue
        subject = f"דוח חודשי מאוחד — מטוס {aircraft} — {month_he} {year}"
        message = f"""
        <div dir="rtl" style="font-family: Arial, sans-serif;">
            <p>שלום,</p>
            <p>מצורף דוח חודשי מאוחד למטוס <strong>{aircraft}</strong> לחודש <strong>{month_he} {year}</strong>.</p>
            <p>הדוח כולל גיליון עבודות וזמנים וגיליון סיכום חלפים.</p>
            <p>בברכה,</p>
            <p>Arrow Aviation MX System</p>
        </div>
        """
        try:
            frappe.sendmail(
                recipients=recipients,
                subject=subject,
                message=message,
                attachments=[
                    {
                        "fname": f"Monthly_Report_{aircraft}_{year}-{month:02d}.xlsx",
                        "fcontent": content,
                    }
                ],
                reference_doctype="Work Report",
            )
            sent.append(aircraft)
        except Exception as e:
            frappe.log_error(f"Monthly report send failed for {aircraft}: {e}", "Monthly Report")
            skipped.append(aircraft)

    return {"year": year, "month": month, "aircraft": sent, "skipped": skipped}
